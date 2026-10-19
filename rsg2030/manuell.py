"""Manuell erfasste Planungsmatrizen für PDF-Seiten ohne Textebene (Bilder).

Welche Matrizen das sind, ergibt sich aus den PDFs (`matrix_extract.hat_textebene`), nicht aus einer
festen Liste. Derzeit: Burgenland BL-Matrix (Anlage 1, S.1) und Kärnten VR 21 (Anhang S.99), je ein
Blatt in `eingaben/2026/manuell.xlsx` (Blattname «<Land> BL» bzw. «<Land> VR<Nr>»). Eine Zeile je Matrixzeile mit den Spalten
``Abschnitt, Jahr, Block, Zeile_ID, Zeile, <Fachspalten…>, gesamt, Quelle``. Werte so, wie sie im PDF
stehen («9,2», «-», «n.v.»); die Fachzellen sind als Text formatiert, damit die Nachkommastellen
(= Rundungstoleranz) erhalten bleiben. Leere Zeilen (alle Fachspalten leer) gelten als noch nicht
erfasst und werden übersprungen.
"""

from __future__ import annotations

from dataclasses import dataclass
from functools import cache
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rsg2030.config import manuell_xlsx
from rsg2030.matrix_extract import ZEILEN_DEFS, einheit_laut_beschriftung, hat_textebene, zellwert
from rsg2030.seitenregister import Rsg2030Quelle, lade_seitenregister

_ID_SPALTEN = {
    "Abschnitt",
    "Jahr",
    "Block",
    "Planungsgröße",
    "Zeile_ID",
    "Zeile_kanonisch",
    "Zeile_PDF_roh",
    "Zeile",
    "Quelle",
    "Parse_Status",
}
_EINHEIT = {rd.zeile_id: rd.einheit for rd in ZEILEN_DEFS}


@dataclass(frozen=True)
class ManuelleMatrix:
    blatt: str
    land: str
    ebene: str
    region_id: str
    region_name: str
    pdf: str
    seite: int
    """1-basiert wie im PDF-Viewer."""


def manuelle_matrizen() -> tuple[ManuelleMatrix, ...]:
    """Alle Matrizen laut seiten.csv, deren (erste) Seite keine Textebene hat."""
    return _manuelle_matrizen(lade_seitenregister())


@cache  # je Register: das Prüfen der Textebene öffnet die PDFs
def _manuelle_matrizen(quellen: tuple[Rsg2030Quelle, ...]) -> tuple[ManuelleMatrix, ...]:
    out: list[ManuelleMatrix] = []
    for src in quellen:
        if not hat_textebene(src.pdf, src.seite_idx):
            out.append(ManuelleMatrix(f"{src.land} BL", src.land, "BL", src.land, "", src.pdf.name, src.seite_idx + 1))
        for vr in src.vr:
            if not hat_textebene(src.pdf, vr.seite_idx):
                out.append(
                    ManuelleMatrix(
                        f"{src.land} VR{vr.region_id}",
                        src.land,
                        "VR",
                        vr.region_id,
                        vr.region_name,
                        src.pdf.name,
                        vr.seite_idx + 1,
                    )
                )
    return tuple(out)


ANLEITUNG = [
    "Manuelle Erfassung RSG 2030 – Planungsmatrizen von Bildseiten (ohne Textebene im PDF)",
    "",
    "Je Blatt eine Matrix; PDF und Seite stehen in der Spalte «Quelle».",
    "Gelb hinterlegte Zeilen sind noch leer = noch nicht erfasst (werden vom Audit übersprungen).",
    "Werte genau so eintippen, wie sie im PDF stehen: 9,2 · 0,0 · - · n.v. · x",
    "  – Nachkommastellen mittippen (12,0 statt 12): daraus ergibt sich die Rundungstoleranz der Prüfungen.",
    "  – Leere Zelle = im PDF leer.",
    "Spalten Abschnitt, Jahr, Block, Zeile_ID nicht ändern; Zeilen nicht löschen.",
    "",
    "Danach:  python -m rsg2030 daten  →  pytest",
]


def _text(v: object) -> str | None:
    """Zellwert → Rohtext; als Zahl eingegebene Werte werden ins österreichische Format gebracht."""
    if v is None:
        return None
    if isinstance(v, bool):
        raise ValueError(f"Unerwarteter Zellwert: {v!r}")
    if isinstance(v, (int, float)):
        return str(v).replace(".", ",")
    t = str(v).strip()
    return t or None


def lese_manuell(m: ManuelleMatrix, path: Path | None = None) -> pd.DataFrame | None:
    """Blatt der Matrix als Text-DataFrame (None, wenn Datei/Blatt fehlt)."""
    path = path or manuell_xlsx()
    if not path.is_file():
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if m.blatt not in wb.sheetnames:
            return None
        rows = list(wb[m.blatt].iter_rows(values_only=True))
    finally:
        wb.close()
    header = [str(h) for h in rows[0] if h is not None]
    data = [[_text(v) for v in r[: len(header)]] for r in rows[1:] if any(v is not None for v in r)]
    return pd.DataFrame(data, columns=header, dtype=object)


def schreibe_manuell(
    blaetter: dict[str, pd.DataFrame], path: Path | None = None, anleitung: list[str] = ANLEITUNG
) -> None:
    """Erfassungs-Workbook schreiben: Fachzellen als Text, leere Zeilen gelb, Kopf/Labels fixiert."""
    path = path or manuell_xlsx()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anleitung"
    for i, line in enumerate(anleitung, start=1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 110

    gelb = PatternFill("solid", fgColor="FFF2CC")
    grau = PatternFill("solid", fgColor="EDEDED")
    for blatt, df in blaetter.items():
        ws = wb.create_sheet(blatt)
        cols = list(df.columns)
        faecher = [c for c in cols if c not in _ID_SPALTEN]
        ws.append(cols)
        for rec in df.itertuples(index=False):
            ws.append([None if (v is None or (isinstance(v, float) and pd.isna(v))) else str(v) for v in rec])
        n = len(df) + 1
        for j, c in enumerate(cols, start=1):
            col = get_column_letter(j)
            head = ws.cell(1, j)
            head.font = Font(bold=True)
            head.alignment = Alignment(text_rotation=90 if c in faecher else 0, vertical="bottom")
            if c in faecher:
                ws.column_dimensions[col].width = 7
                for i in range(2, n + 1):
                    cell = ws.cell(i, j)
                    cell.number_format = "@"
                    cell.alignment = Alignment(horizontal="right")
            else:
                ws.column_dimensions[col].width = {"Zeile": 60, "Zeile_ID": 22, "Quelle": 45}.get(c, 10)
                for i in range(2, n + 1):
                    ws.cell(i, j).fill = grau
        f0 = get_column_letter(cols.index(faecher[0]) + 1)
        f1 = get_column_letter(cols.index(faecher[-1]) + 1)
        ws.conditional_formatting.add(
            f"A2:{get_column_letter(len(cols))}{n}",
            FormulaRule(formula=[f"COUNTA(${f0}2:${f1}2)=0"], fill=gelb),
        )
        ws.freeze_panes = ws.cell(2, cols.index(faecher[0]) + 1)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"
    wb.save(path)


def _einheit(zeile_id: str) -> str:
    if zeile_id.endswith("_max_vaep_prozent"):
        return "Prozent"
    return _EINHEIT.get(zeile_id, "")


def lade_manuell_long(m: ManuelleMatrix) -> list[dict]:
    """Manuelles Blatt → Records im Format von `matrix_extract` (leere Zeilen übersprungen)."""
    df = lese_manuell(m)
    if df is None:
        return []
    faecher = [c for c in df.columns if c not in _ID_SPALTEN]
    datensaetze: list[dict] = []
    fehler: list[str] = []
    for i, rec in df.iterrows():
        cells = {f: rec[f] for f in faecher if isinstance(rec[f], str) and rec[f].strip()}
        if not cells:
            continue
        for fach, raw in cells.items():
            try:
                wert = zellwert(raw.strip(), _einheit(rec["Zeile_ID"]))  # «-» = 0 gilt für SVE wie für ÄAVE
            except ValueError:
                zelle = f"{get_column_letter(df.columns.get_loc(fach) + 1)}{i + 2}"
                fehler.append(f"  {m.blatt} {zelle} ({rec['Zeile_ID']}, {fach}): {raw!r}")
                continue
            datensaetze.append(
                {
                    "Land": m.land,
                    "Ebene": m.ebene,
                    "Region_ID": m.region_id,
                    "Region_Name": m.region_name,
                    "Block": rec["Block"],
                    "Abschnitt": rec["Abschnitt"],
                    "Jahr": int(rec["Jahr"]) if isinstance(rec["Jahr"], str) and rec["Jahr"].strip() else None,
                    "Zeile_ID": rec["Zeile_ID"],
                    "Einheit": einheit_laut_beschriftung(_einheit(rec["Zeile_ID"]), str(rec.get("Zeile", ""))),
                    "Zeile_PDF_roh": "manuell erfasst",
                    "Fach": fach,
                    "Wert": wert,
                    "Roh": raw.strip(),
                    "PDF": m.pdf,
                    "Seite": m.seite,
                    "Quelle": "manuell",
                }
            )
    if fehler:
        raise ValueError(f"Nicht lesbare Werte in {manuell_xlsx().name}:\n" + "\n".join(fehler))
    return datensaetze
