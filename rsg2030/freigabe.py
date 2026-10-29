"""Freigabe durch einen Menschen: was am PDF bestätigt wird, in einer Mappe (`eingaben/<Stand>/freigabe.xlsx`).

Was in die Mappe kommt, folgt aus drei Regeln (docs/PRUEFUNG.md, «Wer prüft was»):

1. Was eine Maschine gegen die Quelle prüfen kann, prüft ein Test bei jedem Lauf – nicht die Mappe (Seitenregister,
   Richtwerte, Regionen, Rechenprüfungen).
2. Jede Zelle wird zweimal unabhängig gelesen (Parser bzw. Abschrift von Hand; KI aus dem Seitenbild). Ein Mensch
   entscheidet nur, wo die Lesungen nicht übereinstimmen:
   - **Zweitlesung**  Zellen, in denen die Lesungen verschieden sind oder die nicht verglichen werden konnten
   - **Bildseiten**   Seiten ohne Textebene: welche Region auf der Seite steht (auf Textseiten prüft das ein Test am
                      Seitentitel)
3. Was keine Lesung ist, sondern ein Urteil, bestätigt ein Mensch je Eintrag – außer, es steht wörtlich so in der
   ÖSG-Vorlage (vorlage.py):
   - **Spalten**      je Land die Fachspalten, die von der Vorlage abweichen (zusätzlich KFO, ohne STR …), und Kopfzeilen
                      mit einer Zuordnung «Kopf → Fachspalte», die nicht schon der Kopf selbst ist – beide Lesungen
                      benutzen dieselbe Zuordnung
   - **Zeilen**       Zeilenbeschriftung im PDF → `Zeile_ID`, wo die Beschriftung nicht die der Vorlage ist
   - **Fußnoten**     jede Fußnote, die nicht wörtlich der ÖSG-Vorlage entspricht, mit ihrer Einordnung (fussnoten.py);
                      was keine Zahl ändern kann, gesammelt je Land
   - **Entscheidungen** jede Korrektur, jede bekannte Abweichung, jede Ableitung

Jede Zeile trägt einen «Prüfstand»: eine Prüfsumme des Inhalts, der bestätigt wird. Ändert sich dieser Inhalt (andere
Fassung einer Verordnung, geänderte Regel, anders gelesener Wert), verfällt die Bestätigung und die Zeile ist wieder offen.
`python -m rsg2030 daten` prüft die Mappe nur; anlegen bzw. auf den Datenstand bringen (Haken bleiben erhalten):

    python -m rsg2030 freigabe
"""

from __future__ import annotations

import argparse
import hashlib
import logging
import re
import sys
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from rsg2030.ableitungen import ABLEITUNGEN, PROZENT_GESAMT_BEGRUENDUNG, PROZENT_GESAMT_REGEL
from rsg2030.checks import lade_bekannte_abweichungen
from rsg2030.config import freigabe_xlsx
from rsg2030.korrekturen import lade_regeln
from rsg2030.manuell import manuelle_matrizen
from rsg2030.matrix_extract import PROZENT_SUFFIX, kompakt, lese_zahl
from rsg2030.protokoll import einrichten
from rsg2030.schema import GESAMT, ZELLE
from rsg2030.specs import KANONISCHE_ZEILEN
from rsg2030.vorlage import kopf_wie_vorlage, spalten_abweichung, zeile_wie_vorlage

if TYPE_CHECKING:
    from rsg2030.pipeline import PipelineErgebnis

log = logging.getLogger(__name__)


JA = "ja"
LEER = "(leer)"
"""So zeigt die Mappe eine leere Zelle; «leer» oder «(leer)» als Wert laut PDF heißt dasselbe."""
EINGABEN = ("bestaetigt", "pdf_wert", "Anmerkung")


def roh_gleich(a: object, b: object) -> bool:
    """PDF-Texte gleichwertig? Zahlen numerisch («2,0(6)» = «2,0»), sonst Token («-», «n.v.», «x»); «leer» = leer."""
    strich = str.maketrans("–—−", "---")  # Gedankenstrich/Minus wie gedruckt = «-»
    fussnote = lambda t: re.sub(r"\(\d+\)$", "", t.strip().translate(strich))  # noqa: E731  «-(7)» = «-»
    leer = lambda t: "" if t.lower() in ("leer", LEER) else t  # noqa: E731
    ta = leer(fussnote(a)) if isinstance(a, str) else ""
    tb = leer(fussnote(b)) if isinstance(b, str) else ""
    try:
        wa, wb = lese_zahl(ta) if ta else None, lese_zahl(tb) if tb else None
    except ValueError:
        return False
    if wa is not None or wb is not None:
        return wa is not None and wb is not None and abs(wa - wb) < 1e-9
    # «n.v» und «n.v.» stehen beide so in PDFs und bedeuten dasselbe; «x(7)» = «x»
    token = lambda t: "x" if t.lower().startswith("x") else t.lower().rstrip(".")  # noqa: E731
    return token(ta) == token(tb)


def _pruefstand(inhalt: str) -> str:
    return hashlib.sha256(inhalt.encode("utf-8")).hexdigest()[:12]


def _ohne_fussnoten(text: object) -> str:
    """Zeilenbeschriftung wie gedruckt, ohne Fußnotenzeichen."""
    t = re.sub(r"\s+", " ", text if isinstance(text, str) else "").strip()
    return re.sub(r"\s*(\(\d+\)|\*(?=\s|\)|$))", "", t).strip()  # «Gastpatient*innen» bleibt


def _label(text: object) -> str:
    """Zeilenbeschriftung ohne Fußnotenzeichen und Jahreszahl – so, wie sie in mehreren Matrizen gleich vorkommt."""
    return re.sub(r"\b20\d\d\b", "20xx", _ohne_fussnoten(text))


def _gesperrt(text: str) -> bool:
    """Gesperrt gesetzt («Ä A V E n ie d e rg …»): viele einzeln stehende Buchstaben."""
    return len(re.findall(r"(?<!\S)\w(?!\S)", text)) >= 3


def _kanonisch(zeile_id: str) -> str:
    return KANONISCHE_ZEILEN.get(zeile_id.removesuffix("_max_vaep_prozent"), zeile_id)


def _beschriftung(zeile_id: str, roh: object, *, mit_jahr: bool = False) -> str:
    """Beschriftung zum Anzeigen: wie im PDF; gesperrt gesetzte mit der lesbaren Bezeichnung und Vermerk. Ohne
    `mit_jahr` steht jede Jahreszahl als 20xx (für Listen über alle Matrizen); an einer einzelnen Seite wie gedruckt."""
    t = _ohne_fussnoten(roh) if mit_jahr else _label(roh)
    return f"{_kanonisch(zeile_id)} [im PDF gesperrt gesetzt]" if _gesperrt(t) else t


def _einheit(p: PipelineErgebnis, zeile_id: pd.Series, land: pd.Series, bloecke: pd.Series) -> pd.Series:
    """Einheit laut Datensatz je Land, Block und Zeile (im Block «haupt» der SVE-Länder ist «spitalsambulant» SVE).
    `land`/`bloecke` dürfen Listen mit «, » sein – dann alle vorkommenden Einheiten."""
    je = p.long.groupby(["Land", "Block", "Zeile_ID"])["Einheit"].first().to_dict()
    nur_zeile = p.long.drop_duplicates("Zeile_ID").set_index("Zeile_ID")["Einheit"].to_dict()
    out = []
    for z, la, bl in zip(zeile_id, land, bloecke, strict=True):
        e = {je.get((x, b, z)) for x in str(la).split(", ") for b in str(bl).split(", ")} - {None}
        out.append(", ".join(sorted(e)) if e else nur_zeile.get(z, ""))
    return pd.Series(out, index=zeile_id.index)


def _bezugszeilen(p: PipelineErgebnis) -> dict[tuple[str, ...], str]:
    """Je Matrix und Stellen-Zeile ihre Beschriftung wie gedruckt – für den Bezug der %-Zeilen."""
    r = p.raw[
        p.raw["Zeile_ID"].str.startswith(("slst_", "planstellen_")) & ~p.raw["Zeile_ID"].str.endswith(PROZENT_SUFFIX)
    ]
    out: dict[tuple[str, ...], str] = {}
    for k, t in zip(r[[*ZELLE[:-1]]].astype(str).itertuples(index=False), r["Zeile_PDF_roh"]):
        if isinstance(t, str) and t.strip() and t != "manuell erfasst":
            out.setdefault(tuple(k), _ohne_fussnoten(t))
    return out


def _mit_bezug(text: str, schluessel: tuple[str, ...], bezug: dict[tuple[str, ...], str]) -> str:
    """«… in %» mit der Stellen-Zeile, zu der die Prozente gehören; `schluessel` = Matrix + Zeile_ID."""
    zeile_id = schluessel[-1]
    if not zeile_id.endswith(PROZENT_SUFFIX):
        return text
    oben = (*schluessel[:-1], zeile_id.removesuffix(PROZENT_SUFFIX))
    return f"{text} – zu: «{bezug.get(oben) or _kanonisch(oben[-1])}»"


def _zahl(w: object) -> str:
    """Wert wie in der Matrix: Komma, ohne überflüssige Nachkommastellen; kein Wert = «(leer)»."""
    return LEER if w is None or pd.isna(w) else f"{float(w):.10g}".replace(".", ",")  # type: ignore[arg-type]


def _seiten(zellen: pd.DataFrame) -> tuple[str, str]:
    """PDF und Seiten der Zellen (je PDF die Seiten aufsteigend)."""
    z = zellen.dropna(subset=["PDF"])
    if not len(z):
        return "", ""
    pdfs = list(dict.fromkeys(z["PDF"]))
    seiten = sorted({int(x) for x in z["Seite"].dropna()})
    if len(seiten) > 6:
        return ", ".join(pdfs), f"{len(seiten)} Seiten ({seiten[0]}–{seiten[-1]})"
    return ", ".join(pdfs), ", ".join(map(str, seiten))


def _wo(r: pd.Series) -> str:
    region = r["Land"] if r["Ebene"] == "BL" else f"{r['Land']} VR {r['Region_ID']}"
    return f"{region} · {r['Block']} · {r['Abschnitt']}"


# --- Blätter -------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Blatt:
    name: str
    zweck: str
    pruefen: str
    spalten: tuple[tuple[str, str, int], ...]
    """(Überschrift, Feld, Breite) der angezeigten Spalten."""
    baue: Callable[[PipelineErgebnis], pd.DataFrame]
    """Zeilen aus dem Datenstand: angezeigte Felder + «Schluessel» + «Inhalt» (was bestätigt wird)."""
    mit_pdf_wert: bool = False
    """Spalte «Wert laut PDF (falls anders)»: weicht ein eingetragener Wert ab, ist das ein Lesefehler (✗)."""


def _spalten(p: PipelineErgebnis) -> pd.DataFrame:
    """Je Land, dessen Fachspalten von der ÖSG-Vorlage abweichen, eine Zeile mit der Abweichung; dazu je verschiedener
    Kopfzeile eines Landes, in der mindestens ein Kopf nicht schon selbst seine Fachspalte ist, eine Zeile."""
    alle = p.koepfe.copy()
    alle["wo"] = [
        f"{'BL' if r.Ebene == 'BL' else 'VR ' + r.Region_ID} S.{r.Seite}, {r.Nr_auf_Seite}. Kopfzeile"
        for r in alle.itertuples()
    ]
    # je Land eine Zeile, wenn seine Matrizen andere Fachspalten haben als die Vorlage (zusätzlich KFO, ohne STR …)
    land = []
    for name, g in alle.groupby("Land", sort=False):
        abweichung = spalten_abweichung({f for faecher in g["Faecher"] for f in faecher})
        if abweichung:
            land.append(
                dict(
                    Land=name,
                    PDF=g["PDF"].iloc[0],
                    Zuordnung=f"Spalten anders als in der ÖSG-Vorlage – {abweichung}",
                    Vorkommen=g["wo"].iloc[0] + (f" (und {len(g) - 1} weitere Kopfzeilen)" if len(g) > 1 else ""),
                    Darunter=g["Darunter"].iloc[0],
                    Schluessel=f"{name}|Vorlage",
                    Inhalt=abweichung,
                )
            )
    k = alle[[not all(map(kopf_wie_vorlage, r.Kopf_PDF, r.Faecher)) for r in alle.itertuples()]].copy()
    k["Kopf"] = k["Kopf_PDF"].map(" · ".join)
    k["Zuordnung"] = [
        " · ".join(t if t == f else f"{t} → {f}" for t, f in zip(r.Kopf_PDF, r.Faecher)) for r in k.itertuples()
    ]
    # eine Zeile je verschiedener Kopfzeile eines Landes, in der Reihenfolge des ersten Vorkommens (Seite für Seite,
    # auf der Seite von oben nach unten)
    df = (
        k.groupby(["Land", "PDF", "Zuordnung"], sort=False)
        .agg(Kopf=("Kopf", "first"), Vorkommen=("wo", "; ".join), Darunter=("Darunter", "first"))
        .reset_index()
    )
    df["Schluessel"] = df["Land"] + "|" + df["Zuordnung"]
    df["Inhalt"] = df["Zuordnung"]
    return pd.concat([pd.DataFrame(land), df], ignore_index=True)


def _zeilen(p: PipelineErgebnis) -> pd.DataFrame:
    """Je Beschriftung im PDF und ihrer Zeile_ID, außer sie steht so in der ÖSG-Vorlage (vorlage.py)."""
    r = p.raw[p.raw["Quelle"].fillna("PDF").eq("PDF")].copy()  # was der Parser zuordnet (vor Korrekturen)
    r["Label"] = [_beschriftung(z, t) for z, t in zip(r["Zeile_ID"], r["Zeile_PDF_roh"])]
    r["Label"] = r["Label"].replace("", "(ohne Beschriftung: Zeile zwischen spitalsambulant und insgesamt)")
    # gleiche Beschriftung in normaler und gesperrter Schreibweise ist eine Zeile der Liste
    r["Text"] = r["Zeile_PDF_roh"].map(lambda t: kompakt(_label(t)))
    r["Matrix"] = r["Land"] + "|" + r["Ebene"] + "|" + r["Region_ID"]
    df = (
        r.groupby(["Abschnitt", "Text", "Zeile_ID"], sort=False)
        .agg(
            Label=("Label", lambda s: min(s, key=lambda x: ("gesperrt" in x, len(x)))),
            Bloecke=("Block", lambda s: ", ".join(dict.fromkeys(s))),
            Laender=("Land", lambda s: ", ".join(dict.fromkeys(s))),
            Matrizen=("Matrix", "nunique"),
        )
        .reset_index()
    )
    df = df[[not zeile_wie_vorlage(a, t, z) for a, t, z in zip(df["Abschnitt"], df["Text"], df["Zeile_ID"])]].copy()
    df["Bedeutung"] = df["Zeile_ID"].map(lambda z: KANONISCHE_ZEILEN.get(z.removesuffix("_max_vaep_prozent"), ""))
    df.loc[df["Zeile_ID"].str.endswith("_max_vaep_prozent"), "Bedeutung"] = (
        "davon maximal als vertragsärztliche Planstellen in % – zu: " + df["Bedeutung"]
    )
    df["Einheit"] = _einheit(p, df["Zeile_ID"], df["Laender"], df["Bloecke"])
    df["Schluessel"] = df["Abschnitt"] + "|" + df["Text"] + "|" + df["Zeile_ID"]
    df["Inhalt"] = df["Zeile_ID"]
    return df


def _bildseiten(p: PipelineErgebnis) -> pd.DataFrame:
    """Je Seite ohne Textebene: welche Region dort steht (für Textseiten prüft das ein Test am Seitentitel). Die Werte
    sind zweimal unabhängig gelesen – von Hand abgetippt und von der KI aus dem Seitenbild – und laufen wie jede
    andere Zelle über das Blatt «Zweitlesung»."""
    bild = {(m.land, m.region_id) for m in manuelle_matrizen()}
    seiten = []
    for q in p.quellen:
        for ebene, rid, name, idx in [("BL", q.land, "", q.seite_idx)] + [
            ("VR", vr.region_id, vr.region_name, vr.seite_idx) for vr in q.vr
        ]:
            if (q.land, rid) in bild:
                matrix = q.land if ebene == "BL" else f"{q.land} VR {rid} {name}"
                seiten.append(
                    dict(
                        Wo=matrix,
                        PDF=q.pdf.name,
                        Seite=idx + 1,
                        Schluessel=f"Seite|{q.land}|{ebene}|{rid}",
                        Inhalt=f"{q.pdf.name}|{idx + 1}|{matrix}",
                    )
                )
    return pd.DataFrame(seiten, columns=["Wo", "PDF", "Seite", "Schluessel", "Inhalt"])


NICHT_VERGLICHEN = "(nicht verglichen: Zeile unvollständig gelesen)"


def _zweitlesung(p: PipelineErgebnis) -> pd.DataFrame:
    """Je Zelle, in der Lesung A (Datensatz) und Zweitlesung (Seitenbild) verschieden sind, und je Zelle, die nicht
    verglichen werden konnte (Zeile der Zweitlesung unvollständig)."""
    v = p.zweitlesung
    teile = [t for t in (v.abweichungen, v.nicht_verglichen.assign(Roh_B=NICHT_VERGLICHEN)) if len(t)]
    if not teile:
        return pd.DataFrame(columns=["Wo", "Zeile", "Fach", "PDF", "Seite", "Roh", "Roh_B", "Schluessel", "Inhalt"])
    d = pd.concat(teile, ignore_index=True)
    d["Wo"] = d.apply(_wo, axis=1)
    d["Zeile"] = [
        _beschriftung(z, a, mit_jahr=True)
        if isinstance(a, str) and a.strip() and a != "manuell erfasst"
        else (b if isinstance(b, str) and b.strip() else _kanonisch(z))
        for z, a, b in zip(d["Zeile_ID"], d["Zeile_PDF_roh"], d["Zeile"])
    ]
    bezug = _bezugszeilen(p)
    d["Zeile"] = [
        _mit_bezug(t, tuple(k), bezug) for t, k in zip(d["Zeile"], d[ZELLE[:-1]].astype(str).itertuples(index=False))
    ]
    d["Seite"] = d["Seite"].astype(int)
    # leere Zellen sichtbar als «(leer)»; der Prüfstand hängt am Inhalt, nicht an der Anzeige
    d["Roh"] = d["Roh_A"].where(d["Roh_A"].ne(""), LEER)
    d["Inhalt"] = d["Roh_A"] + "|" + d["Roh_B"]
    d["Roh_B"] = d["Roh_B"].where(d["Roh_B"].ne(""), LEER)
    d["Schluessel"] = d[ZELLE].astype(str).agg("|".join, axis=1)
    # zur Orientierung, nicht Teil der Bestätigung: Einheit der Zeile und der Wert, wenn eine Regel ihn gesetzt hat
    long = p.long.assign(Region_ID=p.long["Region_ID"].astype(str))
    d["Region_ID"] = d["Region_ID"].astype(str)
    d["Einheit"] = _einheit(p, d["Zeile_ID"], d["Land"], d["Block"])
    geregelt = long[long["Quelle"].isin(["korrigiert", "abgeleitet"])][[*ZELLE, "Wert", "Regel"]]
    d = d.merge(geregelt, on=ZELLE, how="left")
    d["Nach_Regel"] = [f"{_zahl(w)} ({r})" if isinstance(r, str) and r else "" for w, r in zip(d["Wert"], d["Regel"])]
    return d


def _fussnoten(p: PipelineErgebnis) -> pd.DataFrame:
    """Wörtliche ÖSG-Vorlagen auf Textseiten fehlen: Das prüft die Maschine (Regel 1); auf Bildseiten gibt es keinen
    Text zum Abgleichen. Was Werte betrifft oder auf einer Bildseite steht, je verschiedener Fußnote eine Zeile mit allen
    Stellen; der Rest gesammelt je Land (`fussnoten_blatt`)."""
    from rsg2030.fussnoten import MIT_VERWEIS, ist_vorlage

    df = p.fussnoten.copy()
    bild = {(m.land, m.region_id) for m in manuelle_matrizen()}
    auf_bild = pd.Series([(a, b) in bild for a, b in zip(df["Land"], df["Region_ID"])], index=df.index, dtype=bool)
    df = df[auf_bild | ~df["Text"].map(ist_vorlage)].copy()
    df["_rang"] = [
        0 if e in MIT_VERWEIS else (1 if b else 2) for e, b in zip(df["Einordnung"], auf_bild[df.index], strict=True)
    ]
    return fussnoten_blatt(df)


def fussnoten_blatt(df: pd.DataFrame) -> pd.DataFrame:
    """Fußnoten mit Rang (`_FUSSNOTE_PRUEFEN`) → Zeilen des Blatts, nach Gewicht sortiert. Rang 0 und 1 je verschiedener
    Fußnote (über Länder hinweg), Rang 2 eine Zeile je Land mit allen ihren Fußnoten: Sie ändern keine Zahl, geprüft
    wird nur, ob sich darin doch eine Wertänderung versteckt. Ändert sich eine davon, ist die ganze Sammelzeile offen."""
    from rsg2030.fussnoten import norm

    df = df.copy()
    df["_t"] = df["Text"].map(lambda t: kompakt(norm(t)))
    df["wo"] = [
        f"{land} {'BL' if e == 'BL' else 'VR ' + r} S.{s} {z}"
        for land, e, r, s, z in zip(df["Land"], df["Ebene"], df["Region_ID"], df["Seite"], df["Zeichen"])
    ]
    einzeln = df[df["_rang"] < 2]
    out = (
        einzeln.groupby(["_t", "Einordnung", "Verweis"], sort=False)
        .agg(
            Laender=("Land", lambda x: ", ".join(dict.fromkeys(x))),
            Text=("Text", "first"),
            Hinweis_Verz=("Anmerkung", lambda x: " · ".join(dict.fromkeys(v for v in x if v))),
            Vorkommen=("wo", "; ".join),
            _rang=("_rang", "min"),
        )
        .reset_index()
        .sort_values("_rang", kind="stable")
    )
    out["Schluessel"] = out["_t"] + "|" + out["Laender"]
    out["Inhalt"] = out["Text"] + "|" + out["Einordnung"] + "|" + out["Verweis"]
    sammel = []
    for land, g in df[df["_rang"] == 2].groupby("Land", sort=False):
        je_text = g.drop_duplicates(["_t", "Einordnung"])
        sammel.append(
            dict(
                Laender=land,
                Text="\n".join(f"[{e}] {t}" for e, t in zip(je_text["Einordnung"], je_text["Text"])),
                Einordnung=", ".join(sorted(set(je_text["Einordnung"]))),
                Verweis="",
                Hinweis_Verz="",
                Vorkommen=f"{len(je_text)} verschiedene auf {g['Seite'].nunique()} Seiten",
                _rang=2,
                Schluessel=f"Sammel|{land}",
                Inhalt="\n".join(sorted(f"{t}|{e}" for t, e in zip(je_text["_t"], je_text["Einordnung"]))),
            )
        )
    out = pd.concat([out.drop(columns=["_t"]), pd.DataFrame(sammel)], ignore_index=True)
    out["Pruefen"] = out["_rang"].map(_FUSSNOTE_PRUEFEN)
    return out.drop(columns=["_rang"])


_FUSSNOTE_PRUEFEN = {
    0: "genau: Verweis nachsehen",
    1: "genau: Text am Seitenbild",
    2: "überfliegen, je Land",
}
"""Wie gründlich eine Fußnote zu prüfen ist: Was Werte betrifft oder offen ist, wirkt auf die Daten; auf Bildseiten ist
der Text nur einmal gelesen; der Rest (Hinweis, Teilmenge, abgewandelte Vorlage) ändert keine Zahl – dort nur darauf
achten, dass sich keine Wertänderung darin versteckt («davon» oder «zusätzlich»?)."""


_WERTE_ZEIGEN = 6
"""So viele geänderte Zellen zeigt eine Zeile einzeln («AM: (leer) → 260»); bei mehr die Zahl und die ersten davon."""


def _werte(zellen: pd.DataFrame, alt: str, neu: str | None = None) -> str:
    """Zellen als «VR 92 AM: vorher → nachher» (ohne `neu`: der Wert, wie er im Datensatz steht); bei vielen gekürzt."""
    # Block, Abschnitt und Zeile nur nennen, wenn die Regel mehrere davon betrifft
    extra = [c for c in ("Block", "Abschnitt", "Zeile_ID") if zellen[c].nunique() > 1]
    teile = []
    for _, r in zellen.iterrows():
        wo = " ".join(["BL" if r["Ebene"] == "BL" else f"VR {r['Region_ID']}", *(str(r[c]) for c in extra)])
        teile.append(f"{wo} {r['Fach']}: {_zahl(r[alt])}" + (f" → {_zahl(r[neu])}" if neu else ""))
    if len(teile) <= _WERTE_ZEIGEN:
        return "\n".join(teile)
    return f"{len(teile)} Zellen, z. B.\n" + "\n".join(teile[:3])


def _zeile_im_pdf(p: PipelineErgebnis, zellen: pd.DataFrame, bezug: dict[tuple[str, ...], str]) -> str:
    """Beschriftung der betroffenen Zeile(n) im PDF (wie gedruckt); fehlt die Zeile im PDF, ihre Bedeutung."""
    raw = p.raw.assign(Region_ID=p.raw["Region_ID"].astype(str))
    treffer = raw.merge(zellen[ZELLE].astype({"Region_ID": str}).drop_duplicates(), on=ZELLE)
    texte = []
    for k in dict.fromkeys(zellen[ZELLE[:-1]].astype(str).itertuples(index=False, name=None)):
        z = k[-1]
        roh = treffer.loc[treffer["Zeile_ID"].eq(z), "Zeile_PDF_roh"]
        gedruckt = [t for t in roh if isinstance(t, str) and t.strip() and t != "manuell erfasst"]
        if gedruckt:
            texte += [_mit_bezug(_beschriftung(z, t, mit_jahr=True), k, bezug) for t in gedruckt]
        elif roh.eq("manuell erfasst").any():
            texte.append(_mit_bezug(f"{_kanonisch(z)} [Bildseite]", k, bezug))
        else:
            texte.append(_mit_bezug(f"{_kanonisch(z)} [Zeile nicht im PDF]", k, bezug))
    return " · ".join(dict.fromkeys(texte))


def _passt(df: pd.DataFrame, r: pd.Series, spalten: tuple[str, ...]) -> pd.Series:
    """Zellen von `df`, auf die eine Zeile aus korrekturen.csv bzw. bekannte_abweichungen.csv zutrifft («*» = alle)."""
    m = pd.Series(True, index=df.index)
    for c in spalten:
        v = str(r[c])
        if v != "*":
            m &= df[c].astype(str).eq(v)
    return m


def _entscheidungen(p: PipelineErgebnis) -> pd.DataFrame:
    """Jede Korrektur, jede bekannte Abweichung, jede Ableitung – mit PDF, Seite, Beschriftung im PDF und den Werten
    vorher → nachher, damit am PDF nachgesehen werden kann. Bestätigt wird die Regel (Prüfstand), nicht die Anzeige."""
    long = p.long.assign(Region_ID=p.long["Region_ID"].astype(str))
    log = p.korrekturen.assign(Region_ID=p.korrekturen["Region_ID"].astype(str))
    ort = long[[*ZELLE, "PDF", "Seite"]]
    bezug = _bezugszeilen(p)
    zeilen = []
    for nr, (_, r) in enumerate(lade_regeln().iterrows(), start=1):
        zellen = log[log["Regel_Nr"].eq(nr)].merge(ort, on=ZELLE, how="left")
        pdf, seite = _seiten(zellen)
        einheit = (
            _einheit(p, zellen["Zeile_ID"], zellen["Land"], zellen["Block"])
            if len(zellen)
            else _einheit(p, pd.Series([r["Zeile_ID"]]), pd.Series([r["Land"]]), pd.Series([r["Block"]]))
        )
        zeilen.append(
            dict(
                Art=f"Korrektur ({r['Art']})",
                Wo=f"{r['Land']} · {r['Ebene']} {r['Region_ID']} · {r['Block']} · {r['Abschnitt']}",
                PDF=pdf,
                Seite=seite,
                Zeile=_zeile_im_pdf(p, zellen, bezug) if len(zellen) else _kanonisch(r["Zeile_ID"]),
                Einheit=", ".join(dict.fromkeys(einheit)),
                Zeile_ID=r["Zeile_ID"],
                Fach=r["Fach"],
                Regel=r["Regel"],
                Werte=_werte(zellen, "Wert_PDF", "Wert_neu") if len(zellen) else "(keine Zelle geändert)",
                Begruendung=r["Begruendung"],
                Schluessel="K|"
                + "|".join(
                    str(r[c]) for c in ("Land", "Ebene", "Region_ID", "Block", "Abschnitt", "Zeile_ID", "Fach", "Regel")
                ),
                Inhalt=f"{r['Art']}|{r['Befund']}|{r['Begruendung']}",
            )
        )
    for _, r in lade_bekannte_abweichungen().iterrows():
        # Region_ID = Landesname heißt BL-Matrix
        zellen = long[_passt(long, r, ("Land", "Region_ID", "Block", "Abschnitt", "Zeile_ID", "Fach"))]
        pdf, seite = _seiten(zellen)
        zeilen.append(
            dict(
                Art="bekannte Abweichung",
                Wo=f"{r['Land']} · {r['Region_ID']} · {r['Block']} · {r['Abschnitt']}",
                PDF=pdf,
                Seite=seite,
                Zeile=_zeile_im_pdf(p, zellen, bezug) if len(zellen) else _kanonisch(r["Zeile_ID"]),
                Einheit=", ".join(dict.fromkeys(_einheit(p, zellen["Zeile_ID"], zellen["Land"], zellen["Block"]))),
                Zeile_ID=r["Zeile_ID"],
                Fach=r["Fach"],
                Regel=f"Prüfung {r['Check']} bleibt offen (geht nicht auf oder ist nicht prüfbar)",
                Werte=_werte(zellen, "Wert") if len(zellen) else "",
                Begruendung=r["Grund"],
                Schluessel="B|"
                + "|".join(str(r[c]) for c in ("Check", "Land", "Region_ID", "Block", "Abschnitt", "Zeile_ID", "Fach")),
                Inhalt=str(r["Grund"]),
            )
        )
    for abschnitt, ziel, teile, begruendung in ABLEITUNGEN:
        regel = "summe:" + "+".join(teile)
        zellen = long[long["Quelle"].eq("abgeleitet") & long["Abschnitt"].eq(abschnitt) & long["Zeile_ID"].eq(ziel)]
        matrizen = zellen.drop_duplicates(["Land", "Ebene", "Region_ID", "Block"])
        zeilen.append(
            dict(
                Art="Ableitung",
                Wo=f"alle Matrizen ohne diese Zeile im PDF · {abschnitt}",
                PDF=f"{matrizen['PDF'].nunique()} PDFs",
                Seite=f"{len(matrizen)} Matrizen",
                Zeile=f"{_kanonisch(ziel)} [nicht im PDF, Summe der Teilzeilen]",
                Einheit=", ".join(dict.fromkeys(zellen["Einheit"])),
                Zeile_ID=ziel,
                Fach="*",
                Regel=regel,
                Werte=f"{len(zellen)} Zellen in {len(matrizen)} Matrizen ({', '.join(dict.fromkeys(matrizen['Land']))})",
                Begruendung=begruendung,
                Schluessel=f"A|{abschnitt}|{ziel}",
                Inhalt=f"{regel}|{begruendung}",
            )
        )
    zellen = long[long["Quelle"].eq("abgeleitet") & long["Regel"].eq(PROZENT_GESAMT_REGEL)]
    matrizen = zellen.drop_duplicates(["Land", "Ebene", "Region_ID", "Block"])
    zeilen.append(
        dict(
            Art="Ableitung",
            Wo="alle %-Zeilen ohne gesamt-% im PDF · PLAN",
            PDF=f"{matrizen['PDF'].nunique()} PDFs",
            Seite=f"{len(matrizen)} Matrizen",
            Zeile="- davon maximal als vertragsärztliche Planstellen in % [gesamt nicht im PDF, gewichtetes Mittel]",
            Einheit="Prozent",
            Zeile_ID=f"*{PROZENT_SUFFIX}",
            Fach=GESAMT,
            Regel=PROZENT_GESAMT_REGEL,
            Werte=f"{len(zellen)} Zellen in {len(matrizen)} Matrizen ({', '.join(dict.fromkeys(matrizen['Land']))})",
            Begruendung=PROZENT_GESAMT_BEGRUENDUNG,
            Schluessel=f"A|PLAN|*{PROZENT_SUFFIX}|{GESAMT}",
            Inhalt=f"{PROZENT_GESAMT_REGEL}|{PROZENT_GESAMT_BEGRUENDUNG}",
        )
    )
    return pd.DataFrame(zeilen)


_ZELLSPALTEN = (
    ("Wo", "Wo", 34),
    ("Zeile", "Zeile", 48),
    ("Einheit", "Einheit", 15),
    ("Fach", "Fach", 8),
    ("PDF", "PDF", 30),
    ("Seite", "Seite", 6),
    ("Wert im Datensatz", "Roh", 11),
)

BLAETTER: tuple[Blatt, ...] = (
    Blatt(
        "Zweitlesung",
        "Zellen, in denen die Zweitlesung (Seitenbild) etwas anderes liest als der Datensatz oder die sie nicht verglichen "
        "hat",
        "PDF und Seite öffnen: Was steht dort? Stimmt «Wert im Datensatz» (auch «(leer)»): bestätigt = ja, «Wert laut "
        "PDF» frei lassen (dann hat die Zweitlesung verlesen). Stimmt er nicht: bestätigt = ja und den Wert laut PDF eintragen – das ist ein Lesefehler im Datensatz. "
        f"«{NICHT_VERGLICHEN}»: die KI hat die Zeile mit zu wenigen Werten gelesen; den Wert im Datensatz ebenso prüfen. "
        "«Wert im Datensatz» ist der Text, wie er im PDF gelesen wurde; «nach Regel» zeigt nur zur Information, was eine "
        "Korrektur oder Ableitung daraus macht (bestätigt wird das im Blatt «Entscheidungen»).",
        (*_ZELLSPALTEN, ("Zweitlesung", "Roh_B", 11), ("nach Regel (zur Information)", "Nach_Regel", 30)),
        _zweitlesung,
        mit_pdf_wert=True,
    ),
    Blatt(
        "Bildseiten",
        "Seiten ohne Textebene: welche Region dort steht",
        "PDF und Seite öffnen: Ist das die Planungsmatrix dieser Region (Überschrift, Nummer der Versorgungsregion)? "
        "Wenn ja: bestätigt = ja. Die Werte dieser Seiten prüft das Blatt «Zweitlesung» wie auf jeder anderen Seite – "
        "Abschrift von Hand und KI-Lesung sind zwei unabhängige Lesungen.",
        (("Region laut Seitenregister", "Wo", 34), ("PDF", "PDF", 30), ("Seite", "Seite", 6)),
        _bildseiten,
    ),
    Blatt(
        "Spalten",
        "Fachspalten eines Landes, die von der ÖSG-Vorlage abweichen, und Kopfzeilen mit einer Zuordnung, die nicht schon "
        "der Kopf selbst ist (sonst: Maschine)",
        "«Spalten anders als in der ÖSG-Vorlage»: an einer Kopfzeile des Landes nachsehen – hat die Matrix diese "
        "zusätzlichen Spalten, und fehlen diese wirklich? Was eine zusätzliche Spalte bedeutet, steht in "
        "docs/METHODIK.md (KFO: R4). Sonst: Die Kopfzeile steht genau so an jeder unter «Vorkommen» genannten Stelle – eine davon ansehen genügt (die "
        "wievielte Kopfzeile der Seite, und was direkt darunter steht). Stehen die Köpfe in dieser Reihenfolge da, fehlt "
        "keiner, und ist jede Zuordnung «Kopf → Fachspalte» richtig (z. B. «SON → SON/IDB»)? Kopfzeilen, in denen jeder "
        "Kopf ohne Fußnotenzeichen schon seine Fachspalte ist, stehen nicht hier.",
        (
            ("Land", "Land", 11),
            ("Spaltenköpfe → Fachspalte", "Zuordnung", 120),
            ("Vorkommen (Matrix, Seite, die wievielte Kopfzeile)", "Vorkommen", 50),
            ("Darunter steht (beim ersten Vorkommen)", "Darunter", 40),
        ),
        _spalten,
    ),
    Blatt(
        "Zeilen",
        "Zeilenbeschriftung im PDF, die nicht die der ÖSG-Vorlage ist, und die Zeile, der sie zugeordnet ist",
        "Passt die Beschriftung (wie im PDF, ohne Fußnotenzeichen, Jahr als 20xx) zur Bedeutung der Zeile_ID? "
        "Beschriftungen wie in der ÖSG-Vorlage (Tabellenband S.239; «ÄAVE (ggf. SVE)» auch als «ÄAVE» oder «SVE») "
        "stehen nicht hier. Block «haupt» der SVE-Länder: Zählung des Landes (z. B. SVE spitalsambulant), vergleichbare "
        "ÄAVE erst im Block «aequivalent» – die Zeile_ID nennt die Zeile, nicht die Einheit. Beschreibung aller Zeilen: "
        "docs/DATEN.md, «Zeilen».",
        (
            ("Abschnitt", "Abschnitt", 9),
            ("Block", "Bloecke", 14),
            ("Beschriftung im PDF", "Label", 70),
            ("Zeile_ID", "Zeile_ID", 26),
            ("Einheit", "Einheit", 15),
            ("Bedeutung", "Bedeutung", 60),
            ("Länder", "Laender", 30),
            ("Matrizen", "Matrizen", 8),
        ),
        _zeilen,
    ),
    Blatt(
        "Fußnoten",
        "Jede Fußnote der Matrixseiten, die nicht wörtlich der ÖSG-Vorlage entspricht, mit ihrer Einordnung",
        "Passt die Einordnung? Dass der Text vollständig und wie gedruckt dasteht, prüft auf Textseiten ein Test am PDF – "
        "ins PDF nur bei «genau: Text am Seitenbild» (Bildseite, Text nur von Hand gelesen). «Vorlage»: Text "
        "der ÖSG-Vorlage, Bedeutung in docs/DATEN.md · «Hinweis»: ändert keine Zahl und nicht, was sie bedeutet · "
        "«Teilmenge»: nennt einen davon-Wert, der nicht als Zelle übernommen wird · «betrifft Werte»: unter Verweis "
        "nachsehen, ob es dort so umgesetzt ist · «Methodenfrage»: steht in docs/offene_punkte.md. Einverstanden: "
        "bestätigt = ja. Sonst Anmerkung schreiben und leer lassen. Sortiert nach Gewicht (Spalte «Prüfen»): oben, was "
        "Werte betrifft, und Bildseiten (Text nur einmal gelesen) – genau prüfen. Darunter je Land eine Zeile mit allen "
        "Hinweisen und Teilmengen: überfliegen, ob sich eine Wertänderung darin versteckt («davon» oder «zusätzlich»?).",
        (
            ("Prüfen", "Pruefen", 22),
            ("Länder", "Laender", 14),
            ("Text wie gedruckt", "Text", 80),
            ("Einordnung", "Einordnung", 14),
            ("Verweis", "Verweis", 30),
            ("Anmerkung im Verzeichnis", "Hinweis_Verz", 34),
            ("Vorkommen (Seite, Zeichen)", "Vorkommen", 40),
        ),
        _fussnoten,
    ),
    Blatt(
        "Entscheidungen",
        "Jede Korrektur, jede bekannte Abweichung, jede Ableitung",
        "Am PDF nachsehen, nicht nur die Begründung lesen: Stimmen die genannten Seiten und Zahlen, und folgt die "
        "Regel daraus? «PDF», «Seite», «Zeile im PDF» und «Werte» zeigen, welche Zellen die Regel ändert und wie "
        "(«(leer) → 260»: im PDF leer, im Datensatz 260); bei vielen Zellen die Zahl und Beispiele. Einverstanden: "
        "bestätigt = ja. Sonst Anmerkung schreiben und leer lassen.",
        (
            ("Art", "Art", 22),
            ("Wo", "Wo", 38),
            ("PDF", "PDF", 30),
            ("Seite", "Seite", 12),
            ("Zeile im PDF", "Zeile", 48),
            ("Einheit", "Einheit", 15),
            ("Zeile_ID", "Zeile_ID", 22),
            ("Fach", "Fach", 8),
            ("Regel", "Regel", 34),
            ("Werte (vorher → nachher; bei bekannten Abweichungen: Wert im Datensatz)", "Werte", 40),
            ("Begründung", "Begruendung", 90),
        ),
        _entscheidungen,
    ),
)

ANLEITUNG = [
    "Freigabe RSG 2030 – was ein Mensch am PDF bestätigt",
    "",
    "Jede Zelle ist zweimal unabhängig gelesen (Parser bzw. Abschrift; KI aus dem Seitenbild), und bei jedem Lauf rechnen",
    "die Prüfungen Summen und Teile nach. Hier steht nur, was keine Maschine unabhängig prüfen kann:",
    "",
    *[f"  {b.name:<15} {b.zweck}" for b in BLAETTER],
    "",
    "Was zu tun ist, steht in jedem Blatt in der ersten Zeile. Je Zeile prüfen, in «bestätigt» ja eintragen.",
    "Eine Bestätigung ist eine Aussage über das PDF – nicht über eine Begründung. Im Zweifel leer lassen und anmerken.",
    "Ändert sich bestätigter Inhalt (neue Fassung, andere Regel, anders gelesener Wert), wird die Zeile wieder offen.",
    "",
    "Mappe auf den Datenstand bringen (Haken bleiben):  python -m rsg2030 freigabe",
    "Danach:  python -m rsg2030 daten  →  «FREIGEGEBEN», wenn alles bestätigt ist.",
    "Wer wann bestätigt hat, steht in der Git-Historie dieser Datei: nach dem Abhaken committen (eigener Name als Autor).",
    "Warum genau diese Blätter: docs/PRUEFUNG.md, «Wer prüft was».",
]


# --- Abgleich ------------------------------------------------------------------------------------


@dataclass
class BlattStand:
    gesamt: int = 0
    bestaetigt: int = 0
    neu: int = 0
    """Zeilen, die es in der Mappe noch nicht gibt (→ python -m rsg2030 freigabe)."""
    verfallen: int = 0
    """bestätigt, aber der Inhalt hat sich seither geändert."""
    entfallen: int = 0
    abweichend: pd.DataFrame = field(default_factory=pd.DataFrame)
    """bestätigt mit einem anderen «Wert laut PDF» – Lesefehler."""

    @property
    def offen(self) -> int:
        return self.gesamt - self.bestaetigt


@dataclass
class FreigabeStand:
    blaetter: dict[str, BlattStand]
    geschrieben: bool = False

    @property
    def offen(self) -> int:
        return sum(b.offen for b in self.blaetter.values())

    @property
    def nicht_aktuell(self) -> int:
        """Zeilen, die erst `python -m rsg2030 freigabe` in die Mappe bringt (neu, verfallen, entfallen)."""
        return sum(b.neu + b.verfallen + b.entfallen for b in self.blaetter.values())

    @property
    def abweichend(self) -> pd.DataFrame:
        teile = [b.abweichend.assign(Blatt=n) for n, b in self.blaetter.items() if len(b.abweichend)]
        return pd.concat(teile, ignore_index=True) if teile else pd.DataFrame()


def abgleichen(neu: pd.DataFrame, alt: pd.DataFrame | None, blatt: Blatt) -> tuple[pd.DataFrame, BlattStand]:
    """Neue Zeilen + Eingaben aus der Mappe, wo Schlüssel und Prüfstand gleich sind."""
    neu = neu.copy()
    neu["Pruefstand"] = neu["Inhalt"].map(_pruefstand)
    for c in EINGABEN:
        neu[c] = ""
    stand = BlattStand(gesamt=len(neu))
    if alt is not None and len(alt):
        alt = alt.drop_duplicates("Schluessel").set_index("Schluessel")
        da = neu["Schluessel"].isin(alt.index)
        gleich = da & (neu["Pruefstand"] == neu["Schluessel"].map(alt["Pruefstand"]))
        for c in EINGABEN:
            if c in alt:
                neu.loc[gleich, c] = neu.loc[gleich, "Schluessel"].map(alt[c].fillna("")).astype(str)
        war_bestaetigt = neu["Schluessel"].map(alt["bestaetigt"].fillna("").str.strip().str.lower().eq(JA))
        stand.neu = int((~da).sum())
        stand.verfallen = int((da & ~gleich & war_bestaetigt.fillna(False).astype(bool)).sum())
        stand.entfallen = int((~alt.index.isin(neu["Schluessel"])).sum())
    else:
        stand.neu = len(neu)
    ok = neu["bestaetigt"].str.strip().str.lower().eq(JA)
    stand.bestaetigt = int(ok.sum())
    if blatt.mit_pdf_wert:
        # Lesefehler: bestätigt mit einem Wert laut PDF, der nicht dem Datensatz entspricht
        eingetragen = neu["pdf_wert"].str.strip().ne("")
        wie_daten = pd.Series([roh_gleich(a, b) for a, b in zip(neu["pdf_wert"], neu["Roh"])], index=neu.index)
        stand.abweichend = neu.loc[ok & eingetragen & ~wie_daten].reindex(
            columns=[f for _, f, _ in blatt.spalten] + ["pdf_wert"]
        )
    return neu, stand


# --- Excel ---------------------------------------------------------------------------------------

_KOPF_EINGABE = {
    "bestaetigt": ("bestätigt", 10),
    "pdf_wert": ("Wert laut PDF (falls anders)", 14),
    "Anmerkung": ("Anmerkung", 40),
}
_VERSTECKT = ("Schluessel", "Pruefstand")


def _felder(blatt: Blatt) -> list[tuple[str, str, int]]:
    eingaben = ["bestaetigt", *(["pdf_wert"] if blatt.mit_pdf_wert else []), "Anmerkung"]
    return [
        *blatt.spalten,
        *((_KOPF_EINGABE[e][0], e, _KOPF_EINGABE[e][1]) for e in eingaben),
        *((v, v, 14) for v in _VERSTECKT),
    ]


def lese(path: Path | None = None) -> dict[str, pd.DataFrame]:
    """Blätter der Mappe (fehlt sie, leer). Spalten unter ihren Feldnamen."""
    path = path or freigabe_xlsx()
    if not path.is_file():
        return {}
    out: dict[str, pd.DataFrame] = {}
    for blatt in BLAETTER:
        try:
            df = pd.read_excel(path, sheet_name=blatt.name, dtype=str, skiprows=1, keep_default_na=False)
        except ValueError:  # Blatt fehlt
            continue
        out[blatt.name] = df.rename(columns={kopf: f for kopf, f, _ in _felder(blatt)})
    return out


def _ganzzahl(v: object) -> bool:
    try:
        return v is not None and float(str(v)).is_integer()
    except ValueError:  # «35, 36» oder «8 Seiten (3–26)»
        return False


def schreibe(blaetter: dict[str, pd.DataFrame], path: Path | None = None) -> None:
    path = path or freigabe_xlsx()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anleitung"
    for i, t in enumerate(ANLEITUNG, start=1):
        ws.cell(i, 1, t)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 130
    grau, gruen, rot = (PatternFill("solid", fgColor=c) for c in ("EDEDED", "C6EFCE", "F8CBAD"))
    for blatt in BLAETTER:
        df = blaetter[blatt.name]
        felder = _felder(blatt)
        ws = wb.create_sheet(blatt.name)
        ws.append([f"Prüfen: {blatt.pruefen}"])
        ws["A1"].font = Font(italic=True)
        ws.append([kopf for kopf, _, _ in felder])
        for _, r in df.iterrows():
            zeile = []
            for _, f, _ in felder:
                v = r.get(f)
                v = None if v is None or (isinstance(v, float) and pd.isna(v)) or v == "" else v
                zeile.append(
                    int(float(str(v)))
                    if f in ("Seite", "Matrizen") and _ganzzahl(v)
                    else (None if v is None else str(v))
                )
            ws.append(zeile)
        n = len(df) + 2
        for j, (_, f, breite) in enumerate(felder, start=1):
            spalte = get_column_letter(j)
            ws.cell(2, j).font = Font(bold=True)
            ws.cell(2, j).alignment = Alignment(wrap_text=True, vertical="bottom")
            ws.column_dimensions[spalte].width = breite
            ws.column_dimensions[spalte].hidden = f in _VERSTECKT
            for i in range(3, n + 1):
                c = ws.cell(i, j)
                if f in EINGABEN:
                    c.number_format = "@"
                    c.alignment = Alignment(
                        horizontal="center" if f == "bestaetigt" else "left", wrap_text=f == "Anmerkung"
                    )
                else:
                    c.fill = grau
                    c.alignment = Alignment(wrap_text=breite >= 40, vertical="top")
                    if f == "Roh":
                        c.number_format = "@"
                        c.font = Font(bold=True)
        letzte = get_column_letter(len(felder) - len(_VERSTECKT))
        n = max(n, 3)  # ein Blatt ohne Zeilen (alles von der Maschine geprüft) behält Formatierung und Auswahlliste
        b = get_column_letter([f for _, f, _ in felder].index("bestaetigt") + 1)
        if blatt.mit_pdf_wert:
            w = get_column_letter([f for _, f, _ in felder].index("pdf_wert") + 1)
            ws.conditional_formatting.add(f"A3:{letzte}{n}", FormulaRule(formula=[f"LEN(${w}3)>0"], fill=rot))
        ws.conditional_formatting.add(f"A3:{letzte}{n}", FormulaRule(formula=[f'${b}3="{JA}"'], fill=gruen))
        dv = DataValidation(type="list", formula1=f'"{JA}"', allow_blank=True)
        dv.add(f"{b}3:{b}{n}")
        ws.add_data_validation(dv)
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{letzte}{n}"
    wb.save(path)


# --- Einstieg ------------------------------------------------------------------------------------


def _anzeige_anders(neu: dict[str, pd.DataFrame], alt: dict[str, pd.DataFrame]) -> bool:
    """Weicht eine angezeigte Spalte der Mappe vom Datenstand ab (z. B. neue Beschriftung)? Dann neu schreiben."""
    for blatt in BLAETTER:
        if blatt.name not in alt:
            return True
        felder = [f for _, f, _ in blatt.spalten] + ["Schluessel"]
        a = alt[blatt.name].reindex(columns=felder).fillna("").astype(str).reset_index(drop=True)
        n = neu[blatt.name].reindex(columns=felder).fillna("").astype(str).reset_index(drop=True)
        if not a.equals(n):
            return True
    return False


def _anleitung_anders(path: Path) -> bool:
    """Steht oben in einem Blatt ein anderer Prüftext als im Code (oder fehlt das Blatt)? Dann neu schreiben."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return any(b.name not in wb.sheetnames or wb[b.name]["A1"].value != f"Prüfen: {b.pruefen}" for b in BLAETTER)
    finally:
        wb.close()


def stand(p: PipelineErgebnis, *, schreiben: bool = False, path: Path | None = None) -> FreigabeStand:
    """Mappe mit dem Datenstand abgleichen; mit `schreiben` die Mappe auf den Stand bringen (Haken bleiben)."""
    path = path or freigabe_xlsx()
    alt = lese(path)
    neu: dict[str, pd.DataFrame] = {}
    st: dict[str, BlattStand] = {}
    for blatt in BLAETTER:
        neu[blatt.name], st[blatt.name] = abgleichen(blatt.baue(p), alt.get(blatt.name), blatt)
    ergebnis = FreigabeStand(st)
    if schreiben and (
        ergebnis.nicht_aktuell or not path.is_file() or _anzeige_anders(neu, alt) or _anleitung_anders(path)
    ):
        schreibe(neu, path)
        ergebnis.geschrieben = True
    return ergebnis


def zusammenfassung(st: FreigabeStand) -> str:
    return " · ".join(f"{n} {b.bestaetigt}/{b.gesamt}" for n, b in st.blaetter.items())


def main(argv: list[str] | None = None) -> int:
    einrichten()
    argparse.ArgumentParser(
        prog="python -m rsg2030 freigabe", description="Freigabe-Mappe auf den Datenstand bringen"
    ).parse_args(argv)
    from rsg2030.pipeline import baue_pipeline

    st = stand(baue_pipeline(), schreiben=True)
    log.info(f"{freigabe_xlsx().name}: {zusammenfassung(st)} bestätigt" + (" (aktualisiert)" if st.geschrieben else ""))
    for n, b in st.blaetter.items():
        if b.neu or b.verfallen or b.entfallen:
            log.info(f"  {n}: {b.neu} neu, {b.verfallen} Bestätigungen verfallen, {b.entfallen} entfallen")
    for _, r in st.abweichend.iterrows():
        log.warning(
            f"  ✗ {r['Blatt']}: bestätigt mit anderem Wert laut PDF «{r['pdf_wert']}» statt «{r['Roh']}»: {r['Wo']} {r['Zeile']} {r['Fach']}"
        )
    return 1 if len(st.abweichend) else 0


if __name__ == "__main__":
    sys.exit(main())
