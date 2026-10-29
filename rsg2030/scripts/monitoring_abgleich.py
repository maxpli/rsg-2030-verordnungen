"""Abgleich des IST mit dem RSG-Monitoring der GÖG (von Hand aufgerufen, nicht Teil von «alles»).

Vergleicht das IST jeder BL-Matrix (Long-CSV, vergleichbarer Block laut `schema.bl_long`) mit der Kurzfassung des
RSG-Monitorings desselben Datenjahrs, je Teilzeile und Fach, dazu die Spalte «gesamt». Zusätzlich: AM und IM
spitalsambulant im Monitoring von 2022 auf 2023 und die ZAE laut Fußnote («davon … in ZAE», erst ab Datenjahr 2023) –
daran sieht man, wo die ZAE vorher stecken. Dazu ZMK und KFO getrennt: Das Monitoring führt ZMK «inkl. KFO», also
zeigt die Differenz, ob KFO in der gedruckten ZMK-Spalte eines Landes schon steckt, fehlt oder dazugehört. Was das Ergebnis heißt: docs/PRUEFUNG.md, «Abgleich mit dem
RSG-Monitoring». Kein Wert aus dem Monitoring geht in die Daten; eine Abweichung entscheidet ein Mensch.

Die Kurzfassungen stehen im Register `rsg2030/quellen/quellen.csv` mit Verwendet = «Abgleich» (URL, SHA-256); sie
lädt `python -m rsg2030 download` wie jede Quelle, die Pipeline liest sie nicht. Das Datenjahr liest das Skript aus der
Datei («IST ÄAVE 2022»); eine neue Kurzfassung braucht also nur eine Zeile im Register. Vorher muss
`python -m rsg2030 daten` gelaufen sein.

    python -m rsg2030.scripts.monitoring_abgleich

Wiederholen, wenn sich das IST ändert: neue oder geänderte Korrektur im IST, neue Verordnung.
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

from oesg_methodik.pfade import original_dir
from rsg2030.config import LONG_DATEI, ausgabe_dir
from rsg2030.protokoll import einrichten
from rsg2030.quellen import lade_quellen, pruefe
from rsg2030.schema import bl_long

log = logging.getLogger(__name__)

VERWENDET = "Abgleich"  # Kennzeichnung der Kurzfassungen im Register
JAHR = re.compile(r"IST ÄAVE (\d{4})")

# Zeilenbeschriftung im Monitoring (Anfang, ohne Fußnotenzeichen) → Zeile_ID der Long-CSV
ZEILEN = {
    "ÄAVE spitalsambulant": "spitalsambulant",
    "ÄAVE niedergelassene": "niedergelassen",
    "ÄAVE in selbstständigen Ambulatorien (mit Vertrag)": "selbst_amb",
    "ÄAVE in Kassenambulatorien": "kassen_amb",
    "ÄAVE insgesamt": "insgesamt",
    "ÄAVE Wahlärztinnen": "wahl",
    # «besetzte und unbesetzte Planstellen» (Fußnote); Quelle laut Blatt «Anmerkungen u. Quellen»: ÖGK.
    # Die Zeile darunter («Vertragsärztinnen/-ärzte von BVAEB oder SVS ohne ÖGK-Vertrag») zählt Personen ohne
    # ÖGK-Vertrag, nicht Planstellen je Träger – sie passt zu keiner Zeile der Matrix und bleibt draußen.
    "ÖGK-Planstellen": "planstellen_oegk",
}
# Die PLAN-Zeilen heißen wie die IST-Zeilen («ÄAVE spitalsambulant»): ab dem PLAN-Abschnitt nicht weiterlesen.
ENDE = re.compile(r"^(PLAN (ÄAVE|SVE)|RSG[- ]Planungsstand)")

# Spaltenkopf im Monitoring (ohne Fußnotenzeichen wie «(2)») → Fachspalte der Long-CSV, wo der Name anders ist.
# ORTR = OR + UCH + ORTR (METHODIK R4); ZMK(3) ist «inkl. KFO», daher auf unserer Seite ZMK + KFO.
FACH = {"AM/PV": "AM", "SON": "SON/IDB"}
ORTR_TEILE = ("OR", "UCH", "ORTR")
SCHWELLE_PROZENT = 3  # ÄAVE im Monitoring, ab der ein Abstand in Prozent aussagekräftig ist
# Zeilen, die das Monitoring führt und in denen KFO vorkommt: dort muss ZMK + KFO zu «ZMK inkl. KFO» passen
KFO_ZEILEN = ("niedergelassen", "insgesamt", "planstellen_oegk")
ZAE_FUSSNOTE = re.compile(r"spitalsambulant: davon ([\d.,]+) in ZAE")


def fach(kopf: str) -> str:
    k = re.sub(r"\(\d+\)$", "", kopf.strip())
    return FACH.get(k, k)


def kurzfassungen() -> list[Path]:
    """Die Kurzfassungen laut Register; bricht ab, wenn eine fehlt oder nicht die geprüfte Fassung ist."""
    q = lade_quellen()
    q = q[q["Verwendet"] == VERWENDET]
    if probleme := pruefe(q):
        raise RuntimeError("; ".join(probleme))
    return [original_dir() / d for d in q["Datei"]]


def _zahl(v: object) -> float:
    if v is None or v == "n.v.":
        return float("nan")
    if v == "-":
        return 0.0
    try:
        return float(v)  # type: ignore[arg-type]
    except ValueError:
        return float("nan")


def lies_monitoring(pfad: Path) -> pd.DataFrame:
    """IST je Land (Blatt), Zeile und Fach: Spalten Land, Jahr, Zeile_ID, Fach, Monitoring. Gelesen werden die
    ÄAVE-Zeilen und die ÖGK-Planstellen (`ZEILEN`). Das Datenjahr steht über den Zeilen («IST ÄAVE 2022»)."""
    wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    zeilen = []
    for ws in wb.worksheets:
        kopf: list[str] | None = None
        jahr: int | None = None
        gelesen: set[str] = set()
        for r in ws.iter_rows(values_only=True):
            if kopf is None:
                if len(r) > 3 and r[3] == "AM/PV":
                    kopf = [str(c) if c is not None else "" for c in r]
                continue
            text = next((str(c) for c in r[1:3] if isinstance(c, str) and c.strip()), "")
            if m := JAHR.match(text):
                jahr = int(m.group(1))
                continue
            if ENDE.match(text):
                break
            zid = next((z for anfang, z in ZEILEN.items() if text.startswith(anfang)), None)
            if zid is None:
                continue
            if zid in gelesen:  # dieselbe Zeile zweimal: der PLAN-Abschnitt wurde nicht erkannt
                raise RuntimeError(f"{pfad.name}, Blatt {ws.title}: Zeile «{text}» doppelt")
            gelesen.add(zid)
            if jahr is None:
                raise RuntimeError(f"{pfad.name}, Blatt {ws.title}: kein «IST ÄAVE <Jahr>» über den Zeilen")
            werte = {fach(k): _zahl(v) for k, v in zip(kopf[3:], r[3:]) if k}
            # leere Teilspalte = dort kein eigener Ausweis (UCH in der Planstellen-Zeile), nicht «unbekannt»:
            # nur wenn alle drei leer sind, bleibt ORTR leer
            werte["ORTR"] = pd.Series([werte.get(k) for k in ORTR_TEILE]).sum(min_count=1)
            for k, w in werte.items():
                if k not in ORTR_TEILE[:2]:
                    zeilen.append(dict(Land=ws.title, Jahr=jahr, Zeile_ID=zid, Fach=k, Monitoring=w))
        # Die Planstellen stehen zwischen den IST- und den PLAN-ÄAVE: Fehlen sie, hat sich der Aufbau geändert
        # und der PLAN-Abschnitt beginnt woanders, als `ENDE` annimmt.
        if kopf is not None and "planstellen_oegk" not in gelesen:
            raise RuntimeError(f"{pfad.name}, Blatt {ws.title}: keine Zeile «ÖGK-Planstellen»")
    return pd.DataFrame(zeilen)


def zae_fussnoten(pfad: Path) -> dict[str, float]:
    """ZAE je Land laut Fußnote «IST ÄAVE AM spitalsambulant: davon … in ZAE» (gibt es erst ab Datenjahr 2023)."""
    wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    zae = {}
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            for c in r:
                if isinstance(c, str) and (m := ZAE_FUSSNOTE.search(c)):
                    zae[ws.title] = float(m.group(1).replace(",", "."))
    return zae


def lies_verordnungen(long: pd.DataFrame) -> pd.DataFrame:
    """IST je Land (vergleichbarer Block), Zeile und Fach samt «gesamt»: Spalten Land, Jahr, Zeile_ID, Fach,
    Verordnung, Rundung. KFO geht in ZMK ein, wie im Monitoring («ZMK inkl. KFO»). Die Planstellen stehen in allen
    Ländern im Block «haupt»; `bl_long` nimmt Zeilen, die es nur dort gibt, dazu (METHODIK R6)."""
    teile = []
    for land in long["Land"].unique():
        m = bl_long(long, land)
        m = m[(m["Abschnitt"] == "IST") & m["Zeile_ID"].isin(ZEILEN.values())]
        m = m.assign(Fach=m["Fach"].replace({"KFO": "ZMK"}))
        g = m.groupby(["Land", "Jahr", "Zeile_ID", "Fach"], as_index=False).agg(
            Verordnung=("Wert", lambda s: s.sum(min_count=1)), Rundung=("Rundung", "sum")
        )
        teile.append(g)
    return pd.concat(teile, ignore_index=True)


def kfo_vergleich(long: pd.DataFrame, mon: pd.DataFrame) -> pd.DataFrame:
    """ZMK und KFO getrennt neben dem Monitoring, dessen ZMK «inkl. KFO» ist: Spalten ZMK, KFO, Summe, Monitoring,
    Differenz. Steckt KFO schon in der gedruckten ZMK-Spalte, ist die Summe um KFO zu hoch; fehlt KFO, ist sie um
    KFO zu niedrig. Daran hängt die Lesart der KFO-Fußnote in Oberösterreich (METHODIK.md R4). Mit den Ländern ohne
    KFO-Angabe, denn dort fällt das Fehlen genauso auf."""
    teile = []
    for land in long["Land"].unique():
        m = bl_long(long, land)
        m = m[(m["Abschnitt"] == "IST") & m["Zeile_ID"].isin(KFO_ZEILEN) & m["Fach"].isin(("ZMK", "KFO"))]
        p = m.pivot_table(index=["Land", "Jahr", "Zeile_ID"], columns="Fach", values="Wert", dropna=False)
        teile.append(p.reindex(columns=["ZMK", "KFO"]).reset_index())
    v = pd.concat(teile, ignore_index=True)
    v["Jahr"] = v["Jahr"].astype(int)
    v["Summe"] = v[["ZMK", "KFO"]].sum(axis=1, min_count=1)
    v = v.merge(mon[mon["Fach"] == "ZMK"].drop(columns="Fach"), on=["Land", "Jahr", "Zeile_ID"], how="left")
    v["Differenz"] = (v["Summe"] - v["Monitoring"]).round(3)
    v["Zeile_ID"] = pd.Categorical(v["Zeile_ID"], KFO_ZEILEN)
    return v.sort_values(["Zeile_ID", "Jahr", "Land"]).reset_index(drop=True)


def vergleiche(vo: pd.DataFrame, mon: pd.DataFrame) -> pd.DataFrame:
    """Je Zelle beide Werte, Differenz und ob sie gleich sind. Gleich = auf die doppelte Rundung des gedruckten Werts
    (0,1 bei einer Nachkommastelle, 1 bei ganzen Zahlen wie in Wien). Zellen, die auf beiden Seiten 0 oder leer sind,
    fallen weg."""
    v = vo.merge(mon, on=["Land", "Jahr", "Zeile_ID", "Fach"], how="outer")
    v = v.merge(vo[["Land", "Jahr"]].drop_duplicates())  # nur das IST-Jahr jedes Landes, ohne «Österreich gesamt»
    v["Jahr"] = v["Jahr"].astype(int)
    a, b = v["Verordnung"].fillna(0.0), v["Monitoring"].fillna(0.0)
    v = v[(a != 0) | (b != 0)].copy()
    v["Differenz"] = (v["Verordnung"].fillna(0.0) - v["Monitoring"].fillna(0.0)).round(3)
    toleranz = 2 * v["Rundung"].where(v["Rundung"] > 0, 0.05)
    v["gleich"] = v["Differenz"].abs() <= toleranz + 1e-9
    return v.sort_values(["Zeile_ID", "Jahr", "Land", "Fach"]).reset_index(drop=True)


def _fmt(w: float) -> str:
    return "leer" if pd.isna(w) else f"{w:.1f}".replace(".", ",")


def _tabelle(tab: pd.DataFrame) -> None:
    for z in tab.to_string(na_rep="").splitlines():
        log.info("   %s", z)


def bericht(v: pd.DataFrame, mon: pd.DataFrame, zae: dict[tuple[str, int], float], kfo: pd.DataFrame) -> None:
    log.info(
        "Je Zeile und Land: Fächer verglichen, davon gleich, «gesamt» und die Abweichungen (Verordnung / Monitoring)"
    )
    for zid in ZEILEN.values():
        z = v[v["Zeile_ID"] == zid]
        if z.empty:
            continue
        log.info("")
        log.info("── %s", zid)
        for (land, jahr), g in z.groupby(["Land", "Jahr"], sort=False):
            ges, f = g[g["Fach"] == "gesamt"], g[g["Fach"] != "gesamt"]
            abw = f[~f["gleich"]]
            summe = f"gesamt {_fmt(ges['Verordnung'].sum(min_count=1))} / {_fmt(ges['Monitoring'].sum(min_count=1))}"
            zeile = f"{land} {jahr}: {len(f)} Fächer, {len(f) - len(abw)} gleich · {summe}"
            if abw.empty:
                log.info("   ✓ %s", zeile)
            else:
                liste = "; ".join(f"{r.Fach} {_fmt(r.Verordnung)} / {_fmt(r.Monitoring)}" for r in abw.itertuples())
                log.warning("   ✗ %s · %s", zeile, liste)
    n = v[(v["Zeile_ID"] == "niedergelassen") & (v["Monitoring"] > SCHWELLE_PROZENT)].copy()
    n["Prozent"] = ((n["Verordnung"] / n["Monitoring"] - 1) * 100).round(1)
    tab = n.pivot_table(index=["Jahr", "Land"], columns="Fach", values="Prozent", sort=False)
    tab = tab[[c for c in tab.columns if c != "gesamt"] + ["gesamt"]]
    log.info("")
    log.info("── niedergelassen: Verordnung über Monitoring, in %% (Fächer ab %s ÄAVE im Monitoring)", SCHWELLE_PROZENT)
    _tabelle(tab.sort_index())
    log.info("")
    log.info("── ZMK und KFO getrennt gegen «ZMK inkl. KFO» im Monitoring (IST); «leer» = kein KFO ausgewiesen")
    sp = ["ZMK", "KFO", "Summe", "Monitoring", "Differenz"]
    _tabelle(kfo.set_index(["Zeile_ID", "Jahr", "Land"])[sp].map(_fmt))
    s = mon[
        (mon["Zeile_ID"] == "spitalsambulant") & mon["Fach"].isin(["AM", "IM"]) & (mon["Land"] != "Österreich gesamt")
    ]
    tab = s.pivot_table(index="Land", columns=["Fach", "Jahr"], values="Monitoring", sort=False).round(1)
    for jahr in sorted({j for _, j in zae}):
        tab[f"ZAE {jahr}"] = pd.Series({land: w for (land, j), w in zae.items() if j == jahr})
    log.info("")
    log.info(
        "── spitalsambulant im Monitoring je Datenjahr; ZAE laut Fußnote «davon … in ZAE» (erst ab Datenjahr 2023)"
    )
    _tabelle(tab)


def main(argv: list[str] | None = None) -> int:
    einrichten()
    argparse.ArgumentParser(
        prog="python -m rsg2030.scripts.monitoring_abgleich", description="IST gegen das RSG-Monitoring der GÖG"
    ).parse_args(argv)
    long_csv = ausgabe_dir() / LONG_DATEI
    if not long_csv.exists():
        log.error("❌ %s fehlt → python -m rsg2030 daten", long_csv)
        return 1
    long = pd.read_csv(long_csv)
    vo = lies_verordnungen(long)
    try:
        teile = [(lies_monitoring(p), zae_fussnoten(p)) for p in kurzfassungen()]
    except (RuntimeError, OSError) as e:
        log.error("❌ %s → python -m rsg2030 download", e)
        return 1
    mon = pd.concat([m for m, _ in teile])
    zae = {(land, int(m["Jahr"].iloc[0])): w for m, z in teile for land, w in z.items()}
    bericht(vergleiche(vo, mon), mon, zae, kfo_vergleich(long, mon))
    return 0


if __name__ == "__main__":
    sys.exit(main())
